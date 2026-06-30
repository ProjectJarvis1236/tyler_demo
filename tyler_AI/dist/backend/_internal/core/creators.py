from typing import Optional
from docx import Document
from fpdf import FPDF
import pdfplumber
from qdrant_client import QdrantClient
from qdrant_client.http import models as rest
from sentence_transformers import SentenceTransformer

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import os
import httpx
import asyncio
import json
import configs

import win32com.client
import pandas as pd





def create_txt_file(filepath: str, content: str):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

def create_docx_file(filepath: str, content: str):
    doc = Document()
    for line in content.split('\n'):
        doc.add_paragraph(line)
    doc.save(filepath)


def create_pdf_file(filepath: str, content: str):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # ЖЁСТКО указываем путь к шрифту
    font_path = os.path.abspath("DejaVuSans.ttf")
    if not os.path.exists(font_path): raise Exception(f"Шрифт не найден: {font_path}")

    pdf.add_font("DejaVu", "", font_path, uni=True)
    pdf.set_font("DejaVu", size=12)
    for line in content.split("\n"):
        pdf.multi_cell(0, 10, line)

    pdf.output(filepath)