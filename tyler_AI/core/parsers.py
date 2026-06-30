from docx import Document
import pdfplumber
from openpyxl.utils import get_column_letter
import pandas as pd


def parse_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def parse_pdf(path: str) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() for page in pdf.pages if page.extract_text())


def parse_docx(path: str) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


async def excel_smart_parse(filename, sheet_name):
    df = pd.read_excel(filename, sheet_name=sheet_name)
    columns = []

    for i, col in enumerate(df.columns):
        series = df[col]

        col_info = {
            "name": str(col),
            "letter": get_column_letter(i + 1),
            "non_null": int(series.count())
        }

        if pd.api.types.is_numeric_dtype(series):
            col_info["type"] = "number"
            col_info["mean"] = float(series.mean()) if series.count() > 0 else None
            col_info["min"] = float(series.min()) if series.count() > 0 else None
            col_info["max"] = float(series.max()) if series.count() > 0 else None
        else:
            col_info["type"] = "string"
            col_info["unique"] = int(series.nunique())

        columns.append(col_info)

    return {
        "status": "ok",
        "columns": columns,
        "rows": len(df),
    }
